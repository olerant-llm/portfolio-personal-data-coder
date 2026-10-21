"""復元結果の簡易ビューアー（別ウィンドウ）。

役割:
  - 「元に戻す」で作った復元ファイルの中身を、その場で確認できるようにする。
  - 保存の動きは変えない。保存はすでに済んでいる前提で、その結果を表示するだけ。

表示の切り替え:
  - txt / md（文章）… スクロールできる読み取り専用の文章表示（折り返しあり）。
  - csv / xlsx（表）… 格子（ttk.Treeview）で表示。xlsx が複数シートならシートを切り替えられる。
    行が多いときは先頭 500 行までを表示し、残りは保存ファイルで確認するよう案内する。

内容は読み取り専用（書き換え不可）。文字のコピーはできる。
ビューアーが開けなくても、保存そのものは成功として扱う（呼び出し側で保存先を案内する）。
"""

from __future__ import annotations

import csv
import io
import os
import tkinter as tk
from pathlib import Path
from tkinter import ttk

# 表として格子表示するとき、一度に見せる最大行数。
_MAX_TABLE_ROWS = 500
# 文字コードを上から順に試す（BOM 付き → UTF-8 → cp932）。
_ENCODINGS = ("utf-8-sig", "utf-8", "cp932")


def open_viewer(parent: tk.Misc, out_path, kind: str, header_lines) -> tk.Toplevel:
    """復元結果ビューアーを開いて Toplevel を返す。

    parent      : 親ウィンドウ（Tk か Toplevel）
    out_path    : 復元後に保存したファイルのパス
    kind        : "text" | "markdown" | "csv" | "xlsx"（fileio.Document.kind）
    header_lines: 上部に出す案内文（件数・補足など）の行リスト
    """
    return _Viewer(parent, Path(out_path), kind, list(header_lines or []))


# ---------------------------------------------------------------------------
class _Viewer(tk.Toplevel):
    def __init__(self, parent: tk.Misc, out_path: Path, kind: str, header_lines):
        super().__init__(parent)
        self.configure(background="#F4F5F7")
        self.out_path = out_path
        self.title("復元した内容の確認")
        self.geometry("760x560")
        self.minsize(480, 360)

        self._build_header(header_lines)
        ttk.Separator(self, orient="horizontal").pack(fill="x", padx=12)
        body = ttk.Frame(self, padding=12)
        body.pack(fill="both", expand=True)

        try:
            if kind in ("text", "markdown"):
                self._build_text_view(body)
            elif kind == "csv":
                self._build_table_view(body, self._read_csv_sheets())
            elif kind == "xlsx":
                self._build_table_view(body, self._read_xlsx_sheets())
            else:  # 念のため（未知の種別は文章として見せる）
                self._build_text_view(body)
        except Exception as exc:  # noqa: BLE001
            self._build_error_view(body, exc)

    # --- 上部（保存先とボタン） ----------------------------------------
    def _build_header(self, header_lines) -> None:
        head = ttk.Frame(self, padding=(12, 12, 12, 8))
        head.pack(fill="x")

        ttk.Label(head, text="復元した内容", font=("Yu Gothic UI", 14, "bold"), foreground="#1E5899").pack(anchor="w")
        for line in header_lines:
            ttk.Label(head, text=line, foreground="#666677", justify="left", font=("Yu Gothic UI", 10)).pack(anchor="w")

        ttk.Label(
            head,
            text=f"保存先: {self.out_path}",
            justify="left",
            wraplength=720,
        ).pack(anchor="w", pady=(6, 6))

        btns = ttk.Frame(head)
        btns.pack(anchor="w")
        ttk.Button(
            btns, text="保存先のフォルダを開く", command=self._open_folder
        ).pack(side="left")
        ttk.Button(btns, text="閉じる", command=self.destroy).pack(side="left", padx=(8, 0))

    def _open_folder(self) -> None:
        try:
            os.startfile(str(self.out_path.parent))  # type: ignore[attr-defined]
        except Exception:  # noqa: BLE001 - フォルダを開けなくても致命的ではない
            from tkinter import messagebox

            messagebox.showinfo(
                "フォルダを開けませんでした",
                f"手動で開いてください。\n{self.out_path.parent}",
                parent=self,
            )

    # --- 文章表示（txt / md） ------------------------------------------
    def _build_text_view(self, body: tk.Frame) -> None:
        text = self._read_text(self.out_path)

        wrap = ttk.Frame(body)
        wrap.pack(fill="both", expand=True)
        yscroll = ttk.Scrollbar(wrap, orient="vertical")
        widget = tk.Text(
            wrap,
            wrap="word",
            yscrollcommand=yscroll.set,
            padx=8,
            pady=8,
            font=("Yu Gothic UI", 11),
            bg="white",
            fg="#2D323F",
            insertbackground="#2D323F",
            borderwidth=1,
            relief="solid",
            highlightthickness=0,
        )
        yscroll.config(command=widget.yview)
        yscroll.pack(side="right", fill="y")
        widget.pack(side="left", fill="both", expand=True)

        widget.insert("1.0", text)
        _make_readonly(widget)

    # --- 表表示（csv / xlsx） ------------------------------------------
    def _build_table_view(self, body: tk.Frame, sheets) -> None:
        """sheets: list[(シート名, rows)]。rows は list[list[str]]。"""
        if not sheets:
            self._build_error_view(body, Exception("表示できる内容がありませんでした。"))
            return

        note_var = tk.StringVar()

        # 複数シートならプルダウンで切り替え。
        if len(sheets) > 1:
            bar = ttk.Frame(body)
            bar.pack(fill="x", pady=(0, 8))
            ttk.Label(bar, text="シート:").pack(side="left")
            combo = ttk.Combobox(
                bar,
                values=[name for name, _ in sheets],
                state="readonly",
                width=28,
            )
            combo.current(0)
            combo.pack(side="left", padx=(6, 0))

        holder = ttk.Frame(body)
        holder.pack(fill="both", expand=True)
        ttk.Label(body, textvariable=note_var, foreground="#a30").pack(anchor="w", pady=(6, 0))

        def show_sheet(index: int) -> None:
            for child in holder.winfo_children():
                child.destroy()
            _, rows = sheets[index]
            truncated = _fill_treeview(holder, rows)
            if truncated:
                note_var.set(
                    f"行数が多いため先頭 {_MAX_TABLE_ROWS} 行だけ表示しています。"
                    "以降は保存したファイルで確認してください。"
                )
            else:
                note_var.set("")

        if len(sheets) > 1:
            combo.bind("<<ComboboxSelected>>", lambda e: show_sheet(combo.current()))
        show_sheet(0)

    # --- 読み込み失敗時 -------------------------------------------------
    def _build_error_view(self, body: tk.Frame, exc: Exception) -> None:
        for child in body.winfo_children():
            child.destroy()
        ttk.Label(
            body,
            text=(
                "内容の表示中に問題が起きました。\n"
                "保存自体は済んでいます。上の「保存先のフォルダを開く」から確認してください。\n"
                f"（{exc}）"
            ),
            foreground="#a30",
            justify="left",
            wraplength=700,
        ).pack(anchor="w")

    # --- ファイル読み込み ----------------------------------------------
    @staticmethod
    def _read_text(path: Path) -> str:
        raw = path.read_bytes()
        for enc in _ENCODINGS:
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    def _read_csv_sheets(self):
        text = self._read_text(self.out_path)
        rows = list(csv.reader(io.StringIO(text)))
        return [("表", rows)]

    def _read_xlsx_sheets(self):
        from openpyxl import load_workbook

        wb = load_workbook(self.out_path, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if v is None else str(v) for v in row])
            sheets.append((ws.title, rows))
        wb.close()
        return sheets


# ---------------------------------------------------------------------------
def _fill_treeview(holder: tk.Frame, rows) -> bool:
    """holder に格子（Treeview）を作って rows を表示。

    先頭 _MAX_TABLE_ROWS 行までに切り詰めたら True を返す。
    """
    shown = rows[:_MAX_TABLE_ROWS]
    truncated = len(rows) > _MAX_TABLE_ROWS
    col_count = max((len(r) for r in shown), default=1)
    col_ids = [f"c{i}" for i in range(col_count)]

    yscroll = ttk.Scrollbar(holder, orient="vertical")
    xscroll = ttk.Scrollbar(holder, orient="horizontal")
    tree = ttk.Treeview(
        holder,
        columns=col_ids,
        show="headings",
        yscrollcommand=yscroll.set,
        xscrollcommand=xscroll.set,
    )
    yscroll.config(command=tree.yview)
    xscroll.config(command=tree.xview)

    for i, cid in enumerate(col_ids):
        tree.heading(cid, text=str(i + 1))
        tree.column(cid, width=120, minwidth=48, anchor="w", stretch=False)

    for r in shown:
        values = list(r) + [""] * (col_count - len(r))
        tree.insert("", "end", values=values)

    yscroll.pack(side="right", fill="y")
    xscroll.pack(side="bottom", fill="x")
    tree.pack(side="left", fill="both", expand=True)
    return truncated


def _make_readonly(widget: tk.Text) -> None:
    """書き換えは禁止するが、選択とコピーはできるようにする。"""

    def _block(event):
        # Ctrl 併用のコピー / 全選択、および移動キーは通す。
        ctrl = bool(event.state & 0x4)
        if ctrl and event.keysym.lower() in ("c", "a"):
            return None
        if event.keysym in (
            "Left", "Right", "Up", "Down",
            "Home", "End", "Prior", "Next",
        ):
            return None
        return "break"

    widget.bind("<Key>", _block)
