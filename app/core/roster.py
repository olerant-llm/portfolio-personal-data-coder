"""名簿の読み込み・列の推測・ID採番・マスター管理。

- 受理形式: CSV（UTF-8 / cp932 を自動判定）と xlsx。
- 列の対応付け: 先頭行の見出しから「学年/組/番号/氏名」を推測。
- 内部生徒ID: S + 3桁連番（アプリ内だけで使う）。
- 外部コード: 学年・組・番号・氏名から、登録順に依存しない固定規則で生成する。
"""

from __future__ import annotations

import csv
import datetime as _dt
import hashlib
from pathlib import Path

MASTER_VERSION = 4

# 0/O、1/I など、見間違えやすい文字は使わない。
_CODE_ALPHABET = "23456789ABCDEFGHJKLMNPQRSTUVWXYZ"
_CODE_LETTERS = "ABCDEFGHJKLMNPQRSTUVWXYZ"
_BOUNDED_CODE_LENGTH = 4
_STUDENT_CODE_LENGTH = 5
_CODE_NAMESPACE = "hikozinka:v4"

MIN_GRADE = 1
MAX_GRADE = 9
MIN_GROUP = 1
MAX_GROUP = 20
MIN_NUMBER = 1
MAX_NUMBER = 1000

# 生徒コード5文字の内部値は24bit。上位10bitを番号、下位14bitを
# 学年・組・氏名の指紋に使った後、全体を固定規則で並べ替える。
_STUDENT_NAME_BITS = 14
_STUDENT_VALUE_MASK = (1 << 24) - 1

# 見出し推測用のキーワード（左から優先度が高い）。
_HEADER_KEYWORDS = {
    "grade": ["学年", "年"],
    "class": ["組", "クラス", "class"],
    "number": ["出席番号", "番号", "出席", "no", "number", "#"],
    "name": ["氏名", "名前", "生徒名", "name", "児童名"],
}

# 全角数字 → 半角数字
_ZEN2HAN = str.maketrans("０１２３４５６７８９", "0123456789")


class RosterError(Exception):
    """名簿の読み込み・整形に関する利用者向けエラー。"""


# ---------------------------------------------------------------------------
# 読み込み
# ---------------------------------------------------------------------------
def read_table(path: str | Path) -> tuple[list[str], list[list[str]]]:
    """CSV / xlsx を読み、(見出し行, データ行のリスト) を返す。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".csv":
        return _read_csv(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise RosterError(
        "この形式のファイルは読み込めません。CSV か Excel（.xlsx）を選んでください。"
    )


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp932"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise RosterError("CSVファイルの文字コードを判別できませんでした。")

    rows = list(csv.reader(text.splitlines()))
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise RosterError("ファイルが空です。中身を確認してください。")
    headers = [str(c).strip() for c in rows[0]]
    body = [[str(c).strip() for c in r] for r in rows[1:]]
    return headers, body


def _read_xlsx(path: Path) -> tuple[list[str], list[list[str]]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise RosterError(f"Excelファイルを開けませんでした。（{exc}）") from exc
    try:
        ws = wb.active
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append(cells)
    finally:
        wb.close()

    if not rows:
        raise RosterError("ファイルが空です。中身を確認してください。")
    headers = rows[0]
    body = rows[1:]
    return headers, body


# ---------------------------------------------------------------------------
# 列の推測
# ---------------------------------------------------------------------------
def guess_columns(headers: list[str]) -> dict[str, int | None]:
    """見出し行から各項目の列位置を推測する。

    見つからない項目は None。UI 側で利用者に選ばせる想定。
    """
    result: dict[str, int | None] = {k: None for k in _HEADER_KEYWORDS}
    used: set[int] = set()
    lowered = [h.lower() for h in headers]

    for field, keywords in _HEADER_KEYWORDS.items():
        for kw in keywords:
            kw_l = kw.lower()
            found = None
            for idx, head in enumerate(lowered):
                if idx in used:
                    continue
                if kw_l and kw_l in head:
                    found = idx
                    break
            if found is not None:
                result[field] = found
                used.add(found)
                break
    return result


# ---------------------------------------------------------------------------
# 値の整形
# ---------------------------------------------------------------------------
def _to_int(value: str) -> int | None:
    """「３年」「02」などから整数を取り出す。取れなければ None。"""
    if value is None:
        return None
    s = str(value).translate(_ZEN2HAN)
    digits = "".join(ch for ch in s if ch.isdigit())
    if not digits:
        return None
    return int(digits)


def split_name(name: str) -> tuple[str, str]:
    """「姓 名」を全角/半角スペースで分割。分割できなければ (全体, '')。"""
    s = (name or "").strip()
    for sep in ("　", " "):
        if sep in s:
            parts = [p for p in s.split(sep) if p]
            if len(parts) >= 2:
                return parts[0], "".join(parts[1:])
    return s, ""


def normalize_name_for_code(name: str) -> str:
    """コード計算用に氏名の空白差をなくす。漢字などの文字自体は変えない。"""
    return "".join((name or "").replace("　", " ").split())


def _make_external_code(kind: str, value: str, length: int) -> str:
    """同じ入力ならどの名簿でも同じになる、見た目上ランダムなコードを作る。"""
    message = f"{_CODE_NAMESPACE}|{kind}|{value}".encode("utf-8")
    digest = hashlib.sha256(message).digest()
    # 数字だけのコードは実際の番号に見えるため、先頭は必ず英字にする。
    first = _CODE_LETTERS[digest[0] % len(_CODE_LETTERS)]
    rest = "".join(_CODE_ALPHABET[b & 31] for b in digest[1:length])
    return first + rest


def _validated_code_value(label: str, value: int, minimum: int, maximum: int) -> int:
    """コード化できる整数範囲か確認し、利用者向けのエラーを返す。"""
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise RosterError(f"{label}を数字として読み取れませんでした。") from exc
    if not minimum <= parsed <= maximum:
        raise RosterError(f"{label}は{minimum}～{maximum}の範囲で入力してください。")
    return parsed


def _bounded_code_keys() -> list[tuple[str, tuple[int, ...]]]:
    """4文字へ重複なしで割り当てる全1,209候補を固定順で返す。"""
    keys: list[tuple[str, tuple[int, ...]]] = []
    keys.extend(("grade", (grade,)) for grade in range(MIN_GRADE, MAX_GRADE + 1))
    keys.extend(("group", (group,)) for group in range(MIN_GROUP, MAX_GROUP + 1))
    keys.extend(
        ("number", (number,)) for number in range(MIN_NUMBER, MAX_NUMBER + 1)
    )
    keys.extend(
        ("class", (grade, group))
        for grade in range(MIN_GRADE, MAX_GRADE + 1)
        for group in range(MIN_GROUP, MAX_GROUP + 1)
    )
    return keys


_BOUNDED_CODE_MAP: dict[tuple[str, tuple[int, ...]], str] | None = None


def _bounded_code_map() -> dict[tuple[str, tuple[int, ...]], str]:
    """候補全体を固定規則で4文字へ割り当て、種類をまたぐ重複も防ぐ。"""
    global _BOUNDED_CODE_MAP
    if _BOUNDED_CODE_MAP is not None:
        return _BOUNDED_CODE_MAP

    assigned: dict[tuple[str, tuple[int, ...]], str] = {}
    used: set[str] = set()
    for kind, values in _bounded_code_keys():
        source = "|".join(str(value) for value in values)
        attempt = 0
        while True:
            code = _make_external_code(
                f"bounded:{kind}", f"{source}|{attempt}", _BOUNDED_CODE_LENGTH
            )
            if code not in used:
                assigned[(kind, values)] = code
                used.add(code)
                break
            attempt += 1

    _BOUNDED_CODE_MAP = assigned
    return assigned


def _bounded_code(kind: str, *values: int) -> str:
    return _bounded_code_map()[(kind, tuple(values))]


def make_grade_code(grade: int) -> str:
    """学年単独表記を表す、重複しない4文字コードを作る。"""
    grade_i = _validated_code_value("学年", grade, MIN_GRADE, MAX_GRADE)
    return _bounded_code("grade", grade_i)


def make_group_code(class_no: int) -> str:
    """組単独表記を表す、重複しない4文字コードを作る。"""
    class_i = _validated_code_value("組", class_no, MIN_GROUP, MAX_GROUP)
    return _bounded_code("group", class_i)


def make_number_code(number: int) -> str:
    """番号単独表記を表す、重複しない4文字コードを作る。"""
    number_i = _validated_code_value("番号", number, MIN_NUMBER, MAX_NUMBER)
    return _bounded_code("number", number_i)


def make_class_code(grade: int, class_no: int) -> str:
    """実際の学年・組を表さない、重複しない4文字コードを作る。"""
    grade_i = _validated_code_value("学年", grade, MIN_GRADE, MAX_GRADE)
    class_i = _validated_code_value("組", class_no, MIN_GROUP, MAX_GROUP)
    return _bounded_code("class", grade_i, class_i)


def _encode_base32(value: int, length: int) -> str:
    """整数を、見間違えにくい固定長文字列へ変換する。"""
    chars = [""] * length
    for index in range(length - 1, -1, -1):
        chars[index] = _CODE_ALPHABET[value & 31]
        value >>= 5
    return "".join(chars)


def _encode_student_code(value: int) -> str:
    """24bit整数を、先頭が必ず英字の5文字コードへ変換する。"""
    tail_length = _STUDENT_CODE_LENGTH - 1
    tail_bits = tail_length * 5
    first = _CODE_LETTERS[value >> tail_bits]
    return first + _encode_base32(value & ((1 << tail_bits) - 1), tail_length)


def _permute_student_value(value: int) -> int:
    """24bit整数を、元の番号配置が見えないよう可逆な固定規則で混ぜる。"""
    value = (value + 0x12D687) & _STUDENT_VALUE_MASK
    value ^= value >> 11
    value = (value * 0x13A5B7) & _STUDENT_VALUE_MASK
    value ^= (value << 7) & _STUDENT_VALUE_MASK
    value = (value * 0x0F35A9) & _STUDENT_VALUE_MASK
    value ^= value >> 13
    return value & _STUDENT_VALUE_MASK


def make_external_student_id(
    grade: int, class_no: int, number: int, name: str
) -> str:
    """学年・組・番号・氏名から、登録順に依存しない5文字コードを作る。"""
    grade_i = _validated_code_value("学年", grade, MIN_GRADE, MAX_GRADE)
    class_i = _validated_code_value("組", class_no, MIN_GROUP, MAX_GROUP)
    number_i = _validated_code_value("番号", number, MIN_NUMBER, MAX_NUMBER)
    normalized_name = normalize_name_for_code(name)
    name_digest = hashlib.sha256(
        f"{_CODE_NAMESPACE}|student-name|{grade_i}|{class_i}|{normalized_name}".encode(
            "utf-8"
        )
    ).digest()
    name_fingerprint = int.from_bytes(name_digest[:2], "big") & (
        (1 << _STUDENT_NAME_BITS) - 1
    )
    packed = ((number_i - 1) << _STUDENT_NAME_BITS) | name_fingerprint
    return _encode_student_code(_permute_student_value(packed))


def make_student_id(seq: int) -> str:
    """連番から生徒IDを作る（42 → S042）。"""
    return f"S{seq:03d}"


# ---------------------------------------------------------------------------
# マスター管理
# ---------------------------------------------------------------------------
def new_master() -> dict:
    """空のマスターを作る。"""
    return {
        "version": MASTER_VERSION,
        "created": _dt.datetime.now().isoformat(timespec="seconds"),
        "students": [],
        "next_seq": 1,
    }


def ensure_current_master(master: dict) -> bool:
    """開発版の名簿を現行形式へ更新する。旧外部コードの互換性は持たない。"""
    version = int(master.get("version", 1))
    if version >= MASTER_VERSION:
        return False

    master.pop("code_secret", None)
    master["version"] = MASTER_VERSION
    for student in master.get("students", []):
        number = student.get("number")
        if number is None:
            raise RosterError(
                "番号のない登録があります。登録用名簿に番号を入れて作り直してください。"
            )
        student["external_id"] = make_external_student_id(
            student["grade"], student["class"], number, student["name"]
        )
        student["class_code"] = make_class_code(student["grade"], student["class"])
    return True


def build_student_records(
    headers: list[str],
    body: list[list[str]],
    mapping: dict[str, int | None],
    start_seq: int,
) -> list[dict]:
    """データ行から生徒レコードのリストを作る（ID採番込み）。

    start_seq から連番を振る。呼び出し側で next_seq を更新すること。
    """
    if mapping.get("name") is None:
        raise RosterError("「氏名」の列が指定されていません。")
    if mapping.get("number") is None:
        raise RosterError("「番号」の列が指定されていません。")

    name_i = mapping["name"]
    grade_i = mapping.get("grade")
    class_i = mapping.get("class")
    num_i = mapping.get("number")

    records: list[dict] = []
    seq = start_seq
    for line_no, row in enumerate(body, start=2):  # 1行目は見出し
        def cell(i: int | None) -> str:
            if i is None or i >= len(row):
                return ""
            return str(row[i]).strip()

        name = cell(name_i)
        if not name:
            continue  # 氏名が空の行は飛ばす

        grade = _to_int(cell(grade_i)) if grade_i is not None else None
        class_no = _to_int(cell(class_i)) if class_i is not None else None
        number = _to_int(cell(num_i)) if num_i is not None else None

        if grade is None or class_no is None or number is None:
            raise RosterError(
                f"{line_no}行目「{name}」の学年・組・番号を読み取れませんでした。"
                "それぞれに数字が入っているか確認してください。"
            )

        sei, mei = split_name(name)
        student_id = make_student_id(seq)
        records.append(
            {
                "id": student_id,
                "external_id": make_external_student_id(
                    grade, class_no, number, name
                ),
                "grade": grade,
                "class": class_no,
                "number": number,
                "name": name,
                "sei": sei,
                "mei": mei,
                "class_code": make_class_code(grade, class_no),
            }
        )
        seq += 1
    if not records:
        raise RosterError("登録できる生徒が見つかりませんでした。")
    return records


def add_to_master(
    master: dict,
    headers: list[str],
    body: list[list[str]],
    mapping: dict[str, int | None],
) -> list[dict]:
    """既存マスターに生徒を追加する。追加したレコードのリストを返す。

    既存IDは変えず、next_seq から続けて採番する（IDの再利用はしない）。
    """
    start_seq = int(master.get("next_seq", 1))
    candidates = build_student_records(headers, body, mapping, start_seq)
    plan = classify_student_records(master, candidates)
    if plan["conflicts"]:
        first = plan["conflicts"][0]
        incoming = first["incoming"]
        existing = first["existing"]
        raise RosterError(
            f"同じ学年・組・番号に別の氏名があります。"
            f"（{incoming['grade']}年{incoming['class']}組 "
            f"{incoming['number']}番: 「{existing['name']}」と「{incoming['name']}」）"
        )
    added = plan["additions"]
    for offset, student in enumerate(added):
        student["id"] = make_student_id(start_seq + offset)

    # コードは十分長いが、念のため既存分を含めて衝突を検知する。
    all_students = [*master.get("students", []), *added]
    external_ids: dict[str, tuple[int, int, int, str]] = {}
    class_codes: dict[str, tuple[int, int]] = {}
    for student in all_students:
        external_id = f"{student['class_code']}-{student['external_id']}"
        identity = student_identity_key(student)
        previous_identity = external_ids.setdefault(external_id, identity)
        if previous_identity != identity:
            raise RosterError(
                "生徒コードが偶然重なりました。開発元へ連絡してください。"
            )
        class_key = (student["grade"], student["class"])
        previous_class = class_codes.setdefault(student["class_code"], class_key)
        if previous_class != class_key:
            raise RosterError(
                "クラスコードが偶然重なりました。開発元へ連絡してください。"
            )
    master.setdefault("students", []).extend(added)
    master["next_seq"] = start_seq + len(added)
    return added


def student_identity_key(student: dict) -> tuple[int, int, int, str]:
    """同一人物として扱う「学年・組・番号・氏名」の組を返す。"""
    number = student.get("number")
    if number is None:
        raise RosterError("番号のない登録があります。名簿を確認してください。")
    return (
        int(student["grade"]),
        int(student["class"]),
        int(number),
        normalize_name_for_code(student.get("name", "")),
    )


def classify_student_records(master: dict, records: list[dict]) -> dict:
    """読み込んだ行を、新規・登録済み・同じ席番号の競合に分ける。"""
    exact: dict[tuple[int, int, int, str], dict] = {}
    seats: dict[tuple[int, int, int], dict] = {}
    for student in master.get("students", []):
        identity = student_identity_key(student)
        exact.setdefault(identity, student)
        seats.setdefault(identity[:3], student)

    additions: list[dict] = []
    duplicates: list[dict] = []
    conflicts: list[dict] = []
    for student in records:
        identity = student_identity_key(student)
        if identity in exact:
            duplicates.append(student)
            continue
        existing = seats.get(identity[:3])
        if existing is not None:
            conflicts.append({"incoming": student, "existing": existing})
            continue
        additions.append(student)
        exact[identity] = student
        seats[identity[:3]] = student
    return {
        "additions": additions,
        "duplicates": duplicates,
        "conflicts": conflicts,
    }


def filter_students(
    master: dict,
    grade: int | None = None,
    class_no: int | None = None,
    number: int | None = None,
) -> list[dict]:
    """登録済み名簿を学年・組・番号で絞り、名簿順で返す。"""
    found = []
    for student in master.get("students", []):
        if grade is not None and int(student["grade"]) != int(grade):
            continue
        if class_no is not None and int(student["class"]) != int(class_no):
            continue
        if number is not None and int(student["number"]) != int(number):
            continue
        found.append(student)
    return sorted(
        found,
        key=lambda s: (int(s["grade"]), int(s["class"]), int(s["number"]), s["name"]),
    )


def delete_students(master: dict, student_ids: set[str]) -> list[dict]:
    """指定した内部IDの登録を削除する。内部IDの採番位置は巻き戻さない。"""
    wanted = set(student_ids)
    removed = [s for s in master.get("students", []) if s.get("id") in wanted]
    master["students"] = [
        s for s in master.get("students", []) if s.get("id") not in wanted
    ]
    return removed
