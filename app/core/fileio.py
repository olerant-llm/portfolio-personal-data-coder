"""ファイルの読み書き振り分け（txt / csv / xlsx / docx / pdf）。

考え方:
  - ファイルを「文字列のかたまり（segments）」として取り出す。
  - codec で segments を置換する。
  - 元と同じ形で書き戻す（xlsx は書式を保ったまま、セルの値だけ差し替え）。
  - 出力は必ず別ファイル: 元名_coded.拡張子 / 元名_restored.拡張子。
    元ファイルは読み取りのみで書き換えない。

対応:
  - txt … 文字コードを自動判定し、同じ文字コードで書き戻す。
  - csv … 全セルを走査。文字コードは入力に合わせる。
  - xlsx … openpyxl で文字列セルの値だけ置換。書式は保持。
  - docx / pdf … markitdown で Markdown 化してから置換。出力は .md（元の形式には戻らない）。
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

# 文字コードは上から順に試す。BOM の有無は別途判定して保つ。
_ENCODINGS = ("utf-8", "cp932")
_UTF8_BOM = b"\xef\xbb\xbf"

# markitdown で扱う（=元形式には戻さず .md で出す）拡張子
_MARKDOWN_SOURCES = (".docx", ".pdf")


class FileIOError(Exception):
    """ファイルの読み書きに関する利用者向けエラー。"""


class Document:
    """1つのファイルを読み込んだ結果。segments を書き換えて save に渡す。"""

    def __init__(self, path: Path, kind: str):
        self.path = Path(path)
        self.kind = kind  # "text" | "csv" | "xlsx" | "markdown"
        self.segments: list[str] = []
        # 種類ごとの復元用データ
        self.encoding: str | None = None       # text / csv
        self.csv_shape: list[int] | None = None  # csv: 各行のセル数
        self.workbook = None                    # xlsx
        self.cell_refs: list = []               # xlsx: 文字列セルの参照
        self.output_ext: str | None = None      # markdown 変換時の出力拡張子

    @property
    def note(self) -> str | None:
        """UI に出す補足メッセージ（無ければ None）。"""
        if self.kind == "markdown":
            return (
                "この形式は、文字だけを取り出した .md ファイルとして書き出します。\n"
                "（元の見た目や画像は含まれません）"
            )
        return None


# ---------------------------------------------------------------------------
# 読み込み
# ---------------------------------------------------------------------------
def load(path: str | Path) -> Document:
    """ファイルを読み込んで Document を返す。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".txt":
        return _load_text(path)
    if ext == ".csv":
        return _load_csv(path)
    if ext in (".xlsx", ".xlsm"):
        return _load_xlsx(path)
    if ext in _MARKDOWN_SOURCES or ext == ".md":
        return _load_markdown(path)
    raise FileIOError(
        "この形式のファイルには対応していません。\n"
        "テキスト(.txt) / CSV / Excel(.xlsx) / Word(.docx) / PDF を選んでください。"
    )


def _decode_bytes(raw: bytes) -> tuple[str, str]:
    """バイト列を文字コード自動判定で文字列にする。(文字列, 使った文字コード)。

    BOM の有無まで含めて元と同じ形で書き戻せるよう、
    先頭に UTF-8 BOM があるときだけ utf-8-sig を返す。
    """
    if raw.startswith(_UTF8_BOM):
        return raw.decode("utf-8-sig"), "utf-8-sig"
    for enc in _ENCODINGS:
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    raise FileIOError("ファイルの文字コードを判別できませんでした。")


def _load_text(path: Path) -> Document:
    raw = path.read_bytes()
    text, enc = _decode_bytes(raw)
    doc = Document(path, "text")
    doc.encoding = enc
    doc.segments = [text]
    return doc


def _load_csv(path: Path) -> Document:
    raw = path.read_bytes()
    text, enc = _decode_bytes(raw)
    rows = list(csv.reader(io.StringIO(text)))
    doc = Document(path, "csv")
    doc.encoding = enc
    doc.csv_shape = [len(r) for r in rows]
    segments: list[str] = []
    for r in rows:
        segments.extend(r)
    doc.segments = segments
    return doc


def _load_xlsx(path: Path) -> Document:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path)  # 書式を保つため read_only にしない
    except Exception as exc:  # noqa: BLE001
        raise FileIOError(f"Excelファイルを開けませんでした。（{exc}）") from exc

    doc = Document(path, "xlsx")
    doc.workbook = wb
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    doc.cell_refs.append(cell)
                    doc.segments.append(cell.value)
    return doc


def _load_markdown(path: Path) -> Document:
    ext = path.suffix.lower()
    doc = Document(path, "markdown")
    doc.output_ext = ".md"
    if ext == ".md":
        raw = path.read_bytes()
        text, _ = _decode_bytes(raw)
        doc.segments = [text]
        return doc

    try:
        from markitdown import MarkItDown

        result = MarkItDown().convert(str(path))
        text = getattr(result, "text_content", None)
        if text is None:
            text = getattr(result, "markdown", "")
    except Exception as exc:  # noqa: BLE001
        raise FileIOError(
            f"このファイルから文字を取り出せませんでした。（{exc}）"
        ) from exc
    doc.segments = [text]
    return doc


# ---------------------------------------------------------------------------
# 書き込み
# ---------------------------------------------------------------------------
def output_path(path: str | Path, suffix: str, output_ext: str | None = None) -> Path:
    """出力先パスを作る（元名_suffix.拡張子）。"""
    path = Path(path)
    ext = output_ext or path.suffix
    return path.with_name(f"{path.stem}_{suffix}{ext}")


def save(doc: Document, new_segments: list[str], suffix: str) -> Path:
    """置換後の segments を、元と同じ形で別ファイルに書き出す。パスを返す。"""
    if len(new_segments) != len(doc.segments):
        raise FileIOError("内部エラー: 置換前後の項目数が一致しません。")

    out = output_path(doc.path, suffix, doc.output_ext)
    if doc.kind == "text":
        out.write_bytes(new_segments[0].encode(doc.encoding or "utf-8"))
    elif doc.kind == "markdown":
        out.write_bytes(new_segments[0].encode("utf-8"))
    elif doc.kind == "csv":
        _save_csv(doc, new_segments, out)
    elif doc.kind == "xlsx":
        _save_xlsx(doc, new_segments, out)
    else:  # 念のため
        raise FileIOError("内部エラー: 未知のファイル種別です。")
    return out


def _save_csv(doc: Document, new_segments: list[str], out: Path) -> None:
    # segments を元の行構成に戻す
    rows: list[list[str]] = []
    pos = 0
    for width in doc.csv_shape or []:
        rows.append(new_segments[pos : pos + width])
        pos += width
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerows(rows)
    out.write_bytes(buf.getvalue().encode(doc.encoding or "utf-8"))


def _save_xlsx(doc: Document, new_segments: list[str], out: Path) -> None:
    for cell, value in zip(doc.cell_refs, new_segments):
        cell.value = value  # 値だけ差し替え（書式はそのまま）
    doc.workbook.save(out)
